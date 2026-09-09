import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_admin_guide():
    doc = docx.Document()

    # Set Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styling helper
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Document Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("Masjid Ar-Rahman Web Dashboard")
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1B, 0x4D, 0x3E) # Deep Emerald

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(18)
    sub_run = subtitle_p.add_run("Administrator Operations & Field Configuration Guide")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # Intro
    intro_p = doc.add_paragraph(
        "This guide outlines step-by-step instructions for accessing and updating content on the Masjid Ar-Rahman website. "
        "All updates made in this dashboard publish live to the public-facing site immediately upon saving."
    )
    intro_p.paragraph_format.space_after = Pt(14)

    # Helper function for headings
    def add_custom_heading(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1B, 0x4D, 0x3E)
        return h

    # Section 1
    add_custom_heading("1. Accessing the Dashboard")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.add_run("1. Navigate to the website domain with ").font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    r_code = p.add_run("/admin")
    r_code.font.bold = True
    p.add_run(" appended at the end (e.g., ")
    p.add_run("https://masjid-website-tan.vercel.app/admin").font.italic = True
    p.add_run(").")

    p2 = doc.add_paragraph("2. Enter your assigned Username and Password.")
    p2.paragraph_format.space_after = Pt(4)
    p3 = doc.add_paragraph("3. Click Sign in to access the control panel.")
    p3.paragraph_format.space_after = Pt(12)

    # Section 2
    add_custom_heading("2. Global Actions & Navigation Controls")
    
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Control / Button"
    hdr_cells[1].text = "Function & Description"

    # Style table header
    for cell in hdr_cells:
        cell_xml = parse_xml(r'<w:shd {} w:fill="1B4D3E"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(cell_xml)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    actions = [
        ("View website ↗", "Opens the live public site in a new browser tab for immediate preview."),
        ("Download backup", "Exports and downloads a complete JSON copy of all stored site data."),
        ("Sign out", "Logs out of the admin session and returns to the login authentication screen."),
        ("Save and publish changes", "Fixed in the bottom bar; saves all updated field values and refreshes live site content.")
    ]

    for ctrl, desc in actions:
        row_cells = table.add_row().cells
        row_cells[0].text = ctrl
        row_cells[1].text = desc
        for cell in row_cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 3
    add_custom_heading("3. Field Configuration Breakdown")

    fields_data = [
        ("01 Identity and Welcome", [
            ("Masjid name", "The primary headline name displayed across the website header."),
            ("Organization", "The parent organization (e.g., Nigerian Islamic Society of Massachusetts (NISLAM))."),
            ("Tagline", "Sub-header slogan displayed on the home page."),
            ("Homepage introduction", "Main welcome text introducing visitors to the masjid.")
        ]),
        ("02 Prayer Times", [
            ("Prayer-time source", "Dropdown menu to select between Manual Schedule or Automatic calculation."),
            ("Fajr / Dhuhr / Asr / Maghrib / Isha", "Daily prayer timings in 24-hour format (e.g., 05:26, 12:48, 16:12, 19:03, 20:18)."),
            ("Public prayer-time note", "Informational text displayed under the prayer timetable."),
            ("City / Country", "Location coordinates used when Automatic calculation mode is selected."),
            ("Calculation convention", "Mathematical calculation authority selected for automatic prayer times.")
        ]),
        ("03 Contact and Map", [
            ("Street address", "Physical address visible to all visitors."),
            ("Address note", "Optional auxiliary parking or entryway directions."),
            ("Phone / Email", "Direct contact options for community inquiries."),
            ("Google Maps URL", "Direct Google Maps search link driving map navigation buttons.")
        ]),
        ("04 Secure Donations", [
            ("Zakat / Sadaqah / General fund links", "Direct payment processor URLs (Stripe, PayPal, Donorbox)."),
            ("Public giving note", "Charity registration or tax deduction details displayed on the donation block.")
        ]),
        ("05 Social Channels", [
            ("WhatsApp / Instagram / Facebook", "Full public channel URLs including https://.")
        ]),
        ("06 Announcements and Programs", [
            ("Headline & Label", "Title and category tag (e.g., Community update, Friday prayers)."),
            ("Details", "Full body text containing program times, locations, and descriptions."),
            ("Feature on homepage", "Checkbox to pin the event prominently on the main landing page."),
            ("+ Add announcement", "Button to append a new news card.")
        ])
    ]

    for section_title, fields in fields_data:
        h_sub = doc.add_paragraph()
        h_sub.paragraph_format.space_before = Pt(10)
        h_sub.paragraph_format.space_after = Pt(4)
        h_sub.paragraph_format.keep_with_next = True
        r_sub = h_sub.add_run(section_title)
        r_sub.font.bold = True
        r_sub.font.size = Pt(11.5)
        r_sub.font.color.rgb = RGBColor(0x1B, 0x4D, 0x3E)

        for name, desc in fields:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_after = Pt(2)
            r_name = bp.add_run(f"{name}: ")
            r_name.font.bold = True
            bp.add_run(desc)

    doc.save("Masjid_Admin_Guide.docx")

if __name__ == "__main__":
    create_admin_guide()